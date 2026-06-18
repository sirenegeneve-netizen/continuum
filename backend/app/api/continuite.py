from fastapi import APIRouter, HTTPException, Depends, UploadFile, File
from sqlalchemy.orm import Session
from app.db.database import get_db, ConnecteurDB, EnregistrementDB, ModificationDB, CommentaireDB, UtilisateurDB, ServiceDB
from app.api.auth import oauth2_scheme
from app.security import verifier_token
from app.roles import ROLES_ADMIN_PLATEFORME, ROLES_ADMIN_CONTINUITE, ROLES_ECRITURE_CONTINUITE, Role
from pydantic import BaseModel
from typing import Optional, List
from datetime import datetime
import json, csv, io

router = APIRouter()

# ── HELPERS ─────────────────────────────────────────────────────────

def get_current_user(token: str = Depends(oauth2_scheme), db: Session = Depends(get_db)):
    payload = verifier_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="Token invalide")
    user = db.query(UtilisateurDB).filter(UtilisateurDB.email == payload.get("sub")).first()
    if not user:
        raise HTTPException(status_code=401, detail="Utilisateur introuvable")
    return user

def role_de(user):
    return user.role.value if hasattr(user.role, "value") else user.role

def est_admin_plateforme(user):
    return role_de(user) in [r.value for r in ROLES_ADMIN_PLATEFORME]

def peut_administrer_connecteurs(user):
    return role_de(user) in [r.value for r in ROLES_ADMIN_CONTINUITE]

def peut_ecrire_continuite(user):
    return role_de(user) in [r.value for r in ROLES_ECRITURE_CONTINUITE]

def services_accessibles(user, db: Session):
    if peut_administrer_connecteurs(user):
        return [s.id for s in db.query(ServiceDB).all()]
    return [s.id for s in user.services]

def formater_connecteur(c):
    return {
        "id": c.id,
        "nom": c.nom,
        "type": c.type,
        "description": c.description,
        "global_": c.global_,
        "service_ids": [s.id for s in c.services],
        "champs": c.champs or [],
        "mode_import": c.mode_import,
        "actif": c.actif,
        "date_creation": c.date_creation,
        "nb_enregistrements": len(c.enregistrements)
    }

def formater_enregistrement(e):
    return {
        "id": e.id,
        "connecteur_id": e.connecteur_id,
        "donnees": e.donnees or {},
        "date_import": e.date_import,
        "importe_par": e.importe_par,
        "nb_modifications": len(e.modifications),
        "nb_commentaires": len(e.commentaires)
    }

# ── MODELES ─────────────────────────────────────────────────────────

class ConnecteurCreation(BaseModel):
    nom: str
    type: str
    description: Optional[str] = None
    global_: Optional[bool] = False
    service_ids: Optional[List[int]] = []
    champs: List[dict]
    mode_import: Optional[str] = "fichier"

class ModificationChamp(BaseModel):
    champ: str
    valeur: str

class CommentaireCreation(BaseModel):
    texte: str

class EnregistrementManuel(BaseModel):
    donnees: dict

# ── CONNECTEURS ─────────────────────────────────────────────────────

@router.get("/connecteurs")
def liste_connecteurs(db: Session = Depends(get_db), user=Depends(get_current_user)):
    if not peut_ecrire_continuite(user) and role_de(user) != Role.AUDITEUR.value:
        raise HTTPException(status_code=403, detail="Accès refusé à ce module")
    svc_ids = services_accessibles(user, db)
    connecteurs = db.query(ConnecteurDB).filter(ConnecteurDB.actif == True).all()
    result = []
    for c in connecteurs:
        if c.global_ or any(s.id in svc_ids for s in c.services):
            result.append(formater_connecteur(c))
    return result

@router.post("/connecteurs")
def creer_connecteur(data: ConnecteurCreation, db: Session = Depends(get_db), user=Depends(get_current_user)):
    if not peut_administrer_connecteurs(user):
        raise HTTPException(status_code=403, detail="Réservé aux administrateurs et responsables continuité")
    c = ConnecteurDB(
        nom=data.nom, type=data.type, description=data.description,
        global_=data.global_, champs=data.champs, mode_import=data.mode_import
    )
    if not data.global_ and data.service_ids:
        c.services = db.query(ServiceDB).filter(ServiceDB.id.in_(data.service_ids)).all()
    db.add(c); db.commit(); db.refresh(c)
    return formater_connecteur(c)

@router.put("/connecteurs/{connecteur_id}")
def modifier_connecteur(connecteur_id: int, data: ConnecteurCreation, db: Session = Depends(get_db), user=Depends(get_current_user)):
    if not peut_administrer_connecteurs(user):
        raise HTTPException(status_code=403, detail="Réservé aux administrateurs et responsables continuité")
    c = db.query(ConnecteurDB).filter(ConnecteurDB.id == connecteur_id).first()
    if not c:
        raise HTTPException(status_code=404, detail="Connecteur introuvable")
    c.nom = data.nom; c.type = data.type; c.description = data.description
    c.global_ = data.global_; c.champs = data.champs; c.mode_import = data.mode_import
    c.services = [] if data.global_ else db.query(ServiceDB).filter(ServiceDB.id.in_(data.service_ids or [])).all()
    db.commit(); db.refresh(c)
    return formater_connecteur(c)

@router.delete("/connecteurs/{connecteur_id}")
def supprimer_connecteur(connecteur_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    if not peut_administrer_connecteurs(user):
        raise HTTPException(status_code=403, detail="Réservé aux administrateurs et responsables continuité")
    c = db.query(ConnecteurDB).filter(ConnecteurDB.id == connecteur_id).first()
    if not c:
        raise HTTPException(status_code=404, detail="Connecteur introuvable")
    db.delete(c); db.commit()
    return {"message": "Connecteur supprimé"}

# ── IMPORT ──────────────────────────────────────────────────────────

@router.post("/connecteurs/{connecteur_id}/importer-fichier")
async def importer_fichier(
    connecteur_id: int, fichier: UploadFile = File(...),
    db: Session = Depends(get_db), user=Depends(get_current_user)
):
    if not peut_administrer_connecteurs(user):
        raise HTTPException(status_code=403, detail="Réservé aux administrateurs et responsables continuité")
    c = db.query(ConnecteurDB).filter(ConnecteurDB.id == connecteur_id).first()
    if not c:
        raise HTTPException(status_code=404, detail="Connecteur introuvable")

    contenu = await fichier.read()
    nom = fichier.filename.lower()
    lignes = []

    try:
        if nom.endswith(".json"):
            data = json.loads(contenu)
            lignes = data if isinstance(data, list) else [data]

        elif nom.endswith(".csv") or nom.endswith(".tsv"):
            texte = contenu.decode("utf-8-sig")
            if nom.endswith(".tsv"):
                sep = "\t"
            else:
                premiere_ligne = texte.split("\n")[0]
                sep = ";" if premiere_ligne.count(";") > premiere_ligne.count(",") else ","
            reader = csv.DictReader(io.StringIO(texte), delimiter=sep)
            lignes = [dict(row) for row in reader]
            lignes = [{(k or "").strip(): v for k, v in ligne.items()} for ligne in lignes]

        elif nom.endswith(".xlsx") or nom.endswith(".xls"):
            import openpyxl, tempfile, os
            with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
                tmp.write(contenu)
                tmp_path = tmp.name
            wb = openpyxl.load_workbook(tmp_path, data_only=True)
            ws = wb.active
            headers = [str(cell.value or "").strip() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if any(v is not None for v in row):
                    lignes.append({headers[i]: str(v or "") for i, v in enumerate(row)})
            os.unlink(tmp_path)

        elif nom.endswith(".xml"):
            import xmltodict
            data = xmltodict.parse(contenu)
            def extraire_liste(d):
                if isinstance(d, list): return d
                if isinstance(d, dict):
                    for v in d.values():
                        r = extraire_liste(v)
                        if r: return r
                return []
            lignes = extraire_liste(data)
            if not lignes:
                lignes = [data]

        else:
            raise HTTPException(status_code=400, detail="Format non supporté. Utilisez CSV, TSV, JSON, Excel ou XML.")

        if not lignes:
            raise HTTPException(status_code=400, detail="Le fichier ne contient aucune ligne de données.")

        champs_ids = [ch["id"] for ch in (c.champs or [])]
        colonnes_fichier = list(lignes[0].keys())

        manquants = []
        for cid in champs_ids:
            trouve = any(col == cid or col.upper() == cid.upper() for col in colonnes_fichier)
            if not trouve:
                manquants.append(cid)

        if manquants:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Import refusé. Colonnes manquantes dans le fichier : {', '.join(manquants)}. "
                    f"Colonnes attendues : {', '.join(champs_ids)}. "
                    f"Colonnes trouvées dans le fichier : {', '.join(colonnes_fichier)}. "
                    f"Aucune ligne n'a été importée."
                )
            )

        enregistrements_crees = 0
        for ligne in lignes:
            donnees = {}
            for cid in champs_ids:
                val = ""
                for col in colonnes_fichier:
                    if col == cid or col.upper() == cid.upper():
                        val = ligne.get(col, "")
                        break
                donnees[cid] = str(val).strip()
            e = EnregistrementDB(
                connecteur_id=connecteur_id, donnees=donnees,
                importe_par=f"{user.prenom} {user.nom}"
            )
            db.add(e)
            enregistrements_crees += 1

        db.commit()
        return {"message": f"{enregistrements_crees} enregistrement(s) importé(s)"}

    except HTTPException:
        raise
    except Exception as ex:
        raise HTTPException(status_code=400, detail=f"Erreur import : {str(ex)}")

@router.post("/connecteurs/{connecteur_id}/importer-manuel")
def importer_manuel(connecteur_id: int, data: EnregistrementManuel, db: Session = Depends(get_db), user=Depends(get_current_user)):
    if not peut_administrer_connecteurs(user):
        raise HTTPException(status_code=403, detail="Réservé aux administrateurs et responsables continuité")
    c = db.query(ConnecteurDB).filter(ConnecteurDB.id == connecteur_id).first()
    if not c:
        raise HTTPException(status_code=404, detail="Connecteur introuvable")
    e = EnregistrementDB(connecteur_id=connecteur_id, donnees=data.donnees, importe_par=f"{user.prenom} {user.nom}")
    db.add(e); db.commit(); db.refresh(e)
    return formater_enregistrement(e)

# ── ENREGISTREMENTS ─────────────────────────────────────────────────

@router.get("/connecteurs/{connecteur_id}/enregistrements")
def liste_enregistrements(connecteur_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    c = db.query(ConnecteurDB).filter(ConnecteurDB.id == connecteur_id).first()
    if not c:
        raise HTTPException(status_code=404, detail="Connecteur introuvable")
    svc_ids = services_accessibles(user, db)
    if not c.global_ and not any(s.id in svc_ids for s in c.services):
        raise HTTPException(status_code=403, detail="Accès refusé : ce connecteur n'est pas associé à votre service")
    enregistrements = db.query(EnregistrementDB).filter(EnregistrementDB.connecteur_id == connecteur_id).all()
    return [formater_enregistrement(e) for e in enregistrements]

@router.put("/enregistrements/{enregistrement_id}/champ")
def modifier_champ(enregistrement_id: int, data: ModificationChamp, db: Session = Depends(get_db), user=Depends(get_current_user)):
    if not peut_ecrire_continuite(user):
        raise HTTPException(status_code=403, detail="Votre rôle ne permet pas de modifier ces données (lecture seule)")
    e = db.query(EnregistrementDB).filter(EnregistrementDB.id == enregistrement_id).first()
    if not e:
        raise HTTPException(status_code=404, detail="Enregistrement introuvable")
    c = e.connecteur
    svc_ids = services_accessibles(user, db)
    if not c.global_ and not any(s.id in svc_ids for s in c.services):
        raise HTTPException(status_code=403, detail="Accès refusé")

    champ_config = next((ch for ch in (c.champs or []) if ch["id"] == data.champ), None)
    if champ_config and champ_config.get("verrouille"):
        raise HTTPException(status_code=403, detail=f"Le champ '{champ_config.get('label', data.champ)}' est verrouillé et ne peut pas être modifié")

    valeur_avant = (e.donnees or {}).get(data.champ, "")
    nouvelles_donnees = dict(e.donnees or {})
    nouvelles_donnees[data.champ] = data.valeur
    e.donnees = nouvelles_donnees
    modif = ModificationDB(
        enregistrement_id=enregistrement_id, champ=data.champ,
        valeur_avant=valeur_avant, valeur_apres=data.valeur,
        modifie_par=user.email, modifie_par_nom=f"{user.prenom} {user.nom}"
    )
    db.add(modif); db.commit()
    return {"message": "Champ mis à jour"}

@router.get("/enregistrements/{enregistrement_id}/historique")
def historique_enregistrement(enregistrement_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    e = db.query(EnregistrementDB).filter(EnregistrementDB.id == enregistrement_id).first()
    if not e:
        raise HTTPException(status_code=404, detail="Enregistrement introuvable")
    c = e.connecteur
    svc_ids = services_accessibles(user, db)
    if not c.global_ and not any(s.id in svc_ids for s in c.services):
        raise HTTPException(status_code=403, detail="Accès refusé")
    modifications = sorted(e.modifications, key=lambda m: m.date_modification, reverse=True)
    return [{
        "id": m.id, "champ": m.champ, "valeur_avant": m.valeur_avant,
        "valeur_apres": m.valeur_apres, "modifie_par_nom": m.modifie_par_nom, "date": m.date_modification
    } for m in modifications]

# ── COMMENTAIRES ────────────────────────────────────────────────────

@router.get("/enregistrements/{enregistrement_id}/commentaires")
def liste_commentaires(enregistrement_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    e = db.query(EnregistrementDB).filter(EnregistrementDB.id == enregistrement_id).first()
    if not e:
        raise HTTPException(status_code=404, detail="Enregistrement introuvable")
    c = e.connecteur
    svc_ids = services_accessibles(user, db)
    if not c.global_ and not any(s.id in svc_ids for s in c.services):
        raise HTTPException(status_code=403, detail="Accès refusé")
    return [{
        "id": c2.id, "texte": c2.texte, "auteur_nom": c2.auteur_nom, "date": c2.date
    } for c2 in sorted(e.commentaires, key=lambda x: x.date, reverse=True)]

@router.post("/enregistrements/{enregistrement_id}/commentaires")
def ajouter_commentaire(enregistrement_id: int, data: CommentaireCreation, db: Session = Depends(get_db), user=Depends(get_current_user)):
    if not peut_ecrire_continuite(user):
        raise HTTPException(status_code=403, detail="Votre rôle ne permet pas de commenter (lecture seule)")
    e = db.query(EnregistrementDB).filter(EnregistrementDB.id == enregistrement_id).first()
    if not e:
        raise HTTPException(status_code=404, detail="Enregistrement introuvable")
    c = e.connecteur
    svc_ids = services_accessibles(user, db)
    if not c.global_ and not any(s.id in svc_ids for s in c.services):
        raise HTTPException(status_code=403, detail="Accès refusé")
    commentaire = CommentaireDB(
        enregistrement_id=enregistrement_id, texte=data.texte,
        auteur=user.email, auteur_nom=f"{user.prenom} {user.nom}"
    )
    db.add(commentaire); db.commit()
    return {"message": "Commentaire ajouté"}

# ── EXPORT ──────────────────────────────────────────────────────────

@router.get("/connecteurs/{connecteur_id}/exporter")
def exporter_connecteur(connecteur_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    c = db.query(ConnecteurDB).filter(ConnecteurDB.id == connecteur_id).first()
    if not c:
        raise HTTPException(status_code=404, detail="Connecteur introuvable")
    svc_ids = services_accessibles(user, db)
    if not c.global_ and not any(s.id in svc_ids for s in c.services):
        raise HTTPException(status_code=403, detail="Accès refusé")
    enregistrements = db.query(EnregistrementDB).filter(EnregistrementDB.connecteur_id == connecteur_id).all()
    export = []
    for e in enregistrements:
        ligne = dict(e.donnees or {})
        ligne["_date_import"] = str(e.date_import)
        ligne["_importe_par"] = e.importe_par
        ligne["_nb_modifications"] = len(e.modifications)
        ligne["_commentaires"] = " | ".join([f"{cm.auteur_nom}: {cm.texte}" for cm in e.commentaires])
        export.append(ligne)
    return export
